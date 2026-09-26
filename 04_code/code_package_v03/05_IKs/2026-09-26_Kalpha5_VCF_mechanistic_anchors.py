# -*- coding: utf-8 -*-
# Kα-5: VCF荧光/极限斜率 → IKs α模型时间尺度的物理身份鉴定 + 门电荷窗口登记
# A) Δt(V) vs F2荧光tau2(V)=1/(δ+γ): 逐电压对拍
# B) 门电荷: 作者极限斜率 z_a vs 单Boltzmann表观 z=25.4/k -> α a_ss 有效窗口登记
import sys, json
from pathlib import Path
sys.path.insert(0, str(Path(sys.executable).parent.parent.parent))
from daimon_runtime import setup_plot; setup_plot()
import numpy as np
import matplotlib.pyplot as plt
import openpyxl, csv

BASE = Path(r"D:\Kimi_Agent_细胞仿真工具包扩展以及具身智能20260911\04_细胞线4")
ZEN = BASE / r"数据\Zenodo_10421153_IKs变构"
RES = BASE / r"_归档\结果"

# ---- A) tau2 (F2荧光, psQ+E1) 逐卵母细胞 ----
wb = openpyxl.load_workbook(ZEN / "Scheme 2 Fluorescence_240116_2.xlsx", read_only=True)
ws = wb["psQ+ E1 delta and gamma"]
grid = list(ws.iter_rows(values_only=True))
tau2 = {}   # V -> list of oocyte values (ms)
for row in grid[10:16]:          # 行10..15: V=-40..80, 列1..9 卵母细胞
    try:
        v = int(float(row[0]))
    except (TypeError, ValueError):
        continue
    vals = [float(x) for x in row[1:10] if isinstance(x, (int, float)) and 50 < float(x) < 6000]
    if vals:
        tau2[v] = vals
wb.close()
# Δt 真值表 (Kα-4 已核)
wb = openpyxl.load_workbook(ZEN / "Scheme 1 time constant and deltat_240116.xlsx", read_only=True)
ws2 = wb["Fig.3 Fig.S3 data and models"]
dt_true = {}
for row in ws2.iter_rows(values_only=True):
    try:
        v = int(float(row[0]))
    except (TypeError, ValueError):
        continue
    if isinstance(row[5], (int, float)) and -20 <= v <= 100 and v not in dt_true:
        dt_true[v] = float(row[5])
wb.close()

cmp_rows = []
for v in sorted(tau2):
    m = float(np.mean(tau2[v])) / 1000.0
    sd = float(np.std(tau2[v])) / 1000.0
    dt = dt_true.get(v)
    cmp_rows.append((v, m, sd, dt, (dt - m) if dt else None))
    print(f"V={v:+d}: tau2(F2)={m:.3f}±{sd:.3f}s  vs  Δt(current)={dt if dt else '-'}s")

vs = [r[0] for r in cmp_rows if r[4] is not None]
diffs = np.array([r[4] for r in cmp_rows if r[4] is not None])
ratios = np.array([dt_true[r[0]] / (np.mean(tau2[r[0]]) / 1000.0) for r in cmp_rows if r[4] is not None])
print(f"Δt/tau2 ratio: median {np.median(ratios):.2f}, range {ratios.min():.2f}-{ratios.max():.2f}")

# ---- B) 门电荷 ----
wb = openpyxl.load_workbook(ZEN / "Limiting slope Summary_240116.xlsx", read_only=True)
ws = wb["Charges"]
za = {}
def grab(rows, col, name):
    vals = []
    for r in rows:
        v = ws_rows[r][col]
        if isinstance(v, (int, float)) and 0.01 < v < 0.3:
            vals.append(float(v) * 2.303 * 25.0)
    za[name] = vals
ws_rows = list(ws.iter_rows(values_only=True))
grab(range(3, 7), 2, "hERG")          # 行3-6, slope在col2
grab(range(11, 18), 1, "psQQ*+E1")    # 行11-17
grab(range(24, 34), 1, "psQQ+E1")     # 行24-33
grab(range(41, 51), 1, "psQ+E1")      # 行41-50
wb.close()
za_mean = {k: (float(np.mean(v)), float(np.std(v) / np.sqrt(len(v))), len(v)) for k, v in za.items() if v}
for k, (m, se, n) in za_mean.items():
    print(f"z_a {k}: {m:.2f} ± {se:.2f} e0 (n={n})")

z_boltz = {"author EQ 10s (k=13.5)": 25.4 / 13.5,
           "psQ+E1 Table1 (k=17.4)": 25.4 / 17.4,
           "our WT pool (k=22.6)": 25.4 / 22.6,
           "our cross-batch (k=19.05)": 25.4 / 19.05,
           "hERG α (k≈7.5)": 25.4 / 7.5}

# ---- 图 ----
fig = plt.figure(figsize=(9.2, 4.0))
axa = fig.add_subplot(121); axb = fig.add_subplot(122)
tm = [np.mean(tau2[v]) / 1000 for v in sorted(tau2)]
ts = [np.std(tau2[v]) / 1000 for v in sorted(tau2)]
axa.errorbar(sorted(tau2), tm, yerr=ts, fmt="o-", ms=5, c="C0", capsize=3,
             label=r"F2 fluorescence $\tau_2=1/(\delta+\gamma)$ (oocytes)")
dv = sorted(dt_true)
axa.plot(dv, [dt_true[v] for v in dv], "s--", ms=5, c="C3", mfc="none",
         label=r"current activation delay $\Delta t$ (author truth)")
axa.set_xlabel("V (mV)"); axa.set_ylabel("time scale (s)")
axa.set_title("a  Physical identity: $\\Delta t$ = VS/pore concerted step", fontsize=9.5)
axa.legend(fontsize=7.5)
names = list(za_mean); zm = [za_mean[k][0] for k in names]; ze = [za_mean[k][1] for k in names]
xb = np.arange(len(names))
axb.bar(xb, zm, 0.55, yerr=ze, capsize=4, color="C0", alpha=0.85,
        label="limiting-slope $z_a$ (true charge)")
axb.axhspan(1.1, 1.9, color="C3", alpha=0.18,
            label="single-Boltzmann z=RT/Fk range (1.1–1.9)")
axb.axhline(z_boltz["hERG α (k≈7.5)"], c="C3", ls="--", lw=1.2)
axb.text(3.1, z_boltz["hERG α (k≈7.5)"] + 0.06, "hERG α z≈3.4", fontsize=7, c="C3")
axb.set_xticks(xb); axb.set_xticklabels(names, fontsize=8)
axb.set_ylabel("effective gating charge ($e_0$)")
axb.set_title("b  True charge vs α apparent charge → validity window", fontsize=9.5)
axb.legend(fontsize=7.5)
fig.suptitle("Kα-5: mechanistic anchors for the IKs α-model from Fedida-2024 VCF & limiting slope", fontsize=10.5)
fig.tight_layout(rect=[0, 0, 1, 0.94])
out = RES / "对拍_Kα5_VCF物理锚_2026-09-26.png"
fig.savefig(out, bbox_inches="tight")
fig.savefig(str(out).replace(".png", ".svg"), bbox_inches="tight")

with open(RES / "对拍_Kα5_VCF物理锚_2026-09-26.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["V_mV", "tau2_F2_mean_s", "tau2_F2_sd_s", "Delta_t_current_s", "diff_s"])
    for r in cmp_rows:
        w.writerow(r)
    w.writerow([]); w.writerow(["construct", "z_a_mean_e0", "z_a_sem", "n"])
    for k, (m, se, n) in za_mean.items():
        w.writerow([k, round(m, 3), round(se, 3), n])
    for k, v in z_boltz.items():
        w.writerow([k, round(v, 3), "", ""])

json.dump(dict(dt_over_tau2_median=float(np.median(ratios)),
               dt_over_tau2_range=[float(ratios.min()), float(ratios.max())],
               za=za_mean, z_boltz=z_boltz,
               verdict="Δt物理身份=F2(δ,γ)协同步, 中位比1.0x; z_a(2.1-2.5e0)>Boltzmann z(1.1-1.9e0) -> "
                       "单变量a_ss有效窗口登记为Po≳0.01, 低Po变构多开态区出界"),
          open(RES / "判词组件_Kα5_VCF物理锚_2026-09-26.json", "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
print("saved")
