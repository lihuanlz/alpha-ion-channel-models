# -*- coding: utf-8 -*-
# Kα-4: Zenodo 10421153 (Fedida 2024 JGP original-author workbooks) ground-truth cross-check
# Objects: our Kα2 digitized tables (tau_act Fig3B, Dt Fig3C) + WT G-V pool anchors
# Output: cross-check CSV + triptych PNG + verdict JSON
import sys, json
from pathlib import Path
sys.path.insert(0, str(Path(sys.executable).parent.parent.parent))
from daimon_runtime import setup_plot; setup_plot()
import numpy as np
import matplotlib.pyplot as plt
import openpyxl

BASE = Path(r"D:\Kimi_Agent_细胞仿真工具包扩展以及具身智能20260911\04_细胞线4")
ZEN = BASE / r"数据\Zenodo_10421153_IKs变构"
ARC = BASE / r"_归档\α模型"
RES = BASE / r"_归档\结果"
RES.mkdir(exist_ok=True)

# ---- 1. original-author ground-truth tables ----
wb = openpyxl.load_workbook(ZEN / "Scheme 1 time constant and deltat_240116.xlsx", read_only=True)
ws = wb["Extended tau to +180 mV Fig.3B"]
tau_true = {}
for row in ws.iter_rows(values_only=True):
    # scan each row for adjacent (voltage, tau) numeric pairs; skip leading empty/model columns (tau>20 s treated as model values, discarded)
    nums = []
    for v in row:
        try:
            nums.append(float(v))
        except (TypeError, ValueError):
            nums.append(None)
    for a, b in zip(nums, nums[1:]):
        if a is not None and b is not None and -60 <= a <= 180 and a == int(a) and int(a) % 10 == 0 \
                and 0.1 < b < 20000 and a not in tau_true:
            tau_true[int(a)] = b / 1000.0   # s
            break
ws2 = wb["Fig.3 Fig.S3 data and models"]
dt_true, gv_true = {}, {}
for row in ws2.iter_rows(values_only=True):
    try:
        v = int(float(row[0]))
    except (TypeError, ValueError, IndexError):
        continue
    if isinstance(row[5], (int, float)) and -20 <= v <= 100 and v not in dt_true:
        dt_true[v] = float(row[5])           # s
    if isinstance(row[9], (int, float)) and -70 <= v <= 110 and 0 <= row[9] <= 1.1 \
            and isinstance(row[1], (int, float)) and v not in gv_true:
        gv_true[v] = float(row[9])
wb.close()
print(f"truth: tau {len(tau_true)} pts, dt {len(dt_true)} pts, gv {len(gv_true)} pts")

# ---- 2. our digitized tables ----
tau_ours = {}
for ln in open(ARC / "2026-09-15_Kα2_Fedida2024_Fig3B_wtEQ_数字化.csv", encoding="utf-8-sig"):
    p = ln.strip().split(",")
    if p[0].lstrip("-").isdigit():
        tau_ours[int(p[0])] = float(p[1])
dt_ours = {int(k): v for k, v in json.load(open(ARC / "2026-09-15_Kα2_Δt登记.json", encoding="utf-8"))["dt_full_read"].items()}

# ---- 3. cross-check ----
rows = []
for v in sorted(tau_ours):
    if v in tau_true:
        err = (tau_ours[v] - tau_true[v]) / tau_true[v] * 100
        rows.append(("tau_act", v, tau_ours[v], tau_true[v], err))
tau_errs = np.array([r[4] for r in rows])
dt_rows = []
for v in sorted(dt_ours):
    if v in dt_true:
        err = dt_ours[v] - dt_true[v]
        dt_rows.append(("delta_t", v, dt_ours[v], dt_true[v], err * 1000))  # ms
dt_errs = np.array([r[4] for r in dt_rows])

# G-V: Boltzmann fit of the authors' EQ means
Vv = np.array(sorted(gv_true)); Gg = np.array([gv_true[v] for v in Vv])
from scipy.optimize import curve_fit
def boltz(V, Vh, k): return 1.0 / (1.0 + np.exp(-(V - Vh) / k))
popt, _ = curve_fit(boltz, Vv, Gg, p0=[10, 18])
Vh_auth, k_auth = popt
print(f"author EQ G-V (10 s): V1/2={Vh_auth:.1f} mV, k={k_auth:.1f} mV")

ours_pool = {"Kα1 WT池(Chan 8226585)": (25.05, 22.6), "Kα药跨批参照": (32.12, 19.05)}
lit = {"Fedida psQ+E1 (Table 1, 5 s isochronal)": (4.2, None),
       "Fedida psQQ*+E1 (Table 1)": (23.5, None),
       "Fedida EQ*QQ*Q*+E1 (Table 1)": (32.8, None)}

# ---- 4. write CSV ----
import csv
out_csv = RES / "compare_Ka4_Zenodo10421153_truth_2026-09-26.csv"
with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["quantity", "V_mV", "ours_digitised", "author_truth", "err_%_or_ms"])
    for r in rows + dt_rows:
        w.writerow(r)
    w.writerow([]); w.writerow(["GV_boltzmann", "V1/2_mV", "k_mV", "source"])
    w.writerow(["author_EQ_10s", round(Vh_auth, 2), round(k_auth, 2), "Zenodo10421153 Fig3 sheet"])
    for nm, (vh, k) in {**ours_pool, **lit}.items():
        w.writerow([nm, vh, k if k else "", "sealed/Table1"])

# ---- 5. figure ----
fig = plt.figure(figsize=(13.2, 4.2))
axa = fig.add_subplot(131); axb = fig.add_subplot(132); axc = fig.add_subplot(133)
tv = sorted(tau_true)
axa.plot(tv, [tau_true[v] for v in tv], "o-", ms=4, c="0.2", label="author truth (Zenodo 10421153)")
ov = sorted(tau_ours)
axa.plot(ov, [tau_ours[v] for v in ov], "s--", ms=4, c="C3", mfc="none", label="our Kα2 digitisation")
axa.set_xlabel("V (mV)"); axa.set_ylabel(r"$\tau_{act}$ (s)")
axa.set_title(f"a  τact: max err {np.max(np.abs(tau_errs)):.1f}% (excl. +180 mV: "
              f"{np.max(np.abs(tau_errs[:-1])):.1f}%)", fontsize=9)
axa.legend(fontsize=7)
dv = sorted(dt_true)
axb.plot(dv, [dt_true[v] for v in dv], "o-", ms=4, c="0.2", label="author truth")
od = sorted(dt_ours)
axb.plot(od, [dt_ours[v] for v in od], "s--", ms=4, c="C3", mfc="none", label="our Kα2 digitisation")
axb.set_xlabel("V (mV)"); axb.set_ylabel(r"$\Delta t$ (s)")
axb.set_title(f"b  Δt: systematic low by {np.mean(dt_errs):.0f} ms (max {np.max(np.abs(dt_errs)):.0f} ms)", fontsize=9)
axb.legend(fontsize=7)
Vf = np.linspace(-80, 110, 200)
axc.plot(Vv, Gg, "o", ms=4, c="0.2", label="author EQ 10 s data")
axc.plot(Vf, boltz(Vf, *popt), c="0.2", lw=1,
         label=f"fit: V½={Vh_auth:.1f}, k={k_auth:.1f}")
for nm, (vh, k), c in [("our WT pool 25.05", (25.05, 22.6), "C0"),
                        ("cross-batch 32.12", (32.12, 19.05), "C1")]:
    axc.plot(Vf, boltz(Vf, vh, k), "--", c=c, lw=1.2, label=f"{nm} mV")
    axc.axvline(vh, c=c, lw=0.6, ls=":")
for vh, nm in [(4.2, "psQ+E1 4.2"), (23.5, "psQQ*+E1 23.5"), (32.8, "EQ*QQ*Q* 32.8")]:
    axc.axvline(vh, c="0.55", lw=0.8, ls="-.")
    axc.text(vh + 1, 0.06, nm, fontsize=6.5, rotation=90, va="bottom")
axc.set_xlabel("V (mV)"); axc.set_ylabel("G/Gmax")
axc.set_title("c  G-V: our pool vs author constructs", fontsize=9)
axc.legend(fontsize=6.5, loc="center right")
fig.suptitle("Kα-4: digitised Fedida-2024 tables vs author ground truth (Zenodo 10421153)", fontsize=10.5)
fig.tight_layout(rect=[0, 0, 1, 0.94])
out_png = RES / "compare_Ka4_Zenodo10421153_truth_2026-09-26.png"
fig.savefig(out_png, bbox_inches="tight")
fig.savefig(str(out_png).replace(".png", ".svg"), bbox_inches="tight")

verdict = dict(
    tau_rel_err_pct=dict(max=float(np.max(np.abs(tau_errs))),
                         max_excl_180=float(np.max(np.abs(tau_errs[:-1]))),
                         rmse=float(np.sqrt(np.mean(tau_errs ** 2)))),
    dt_err_ms=dict(mean=float(np.mean(dt_errs)), max_abs=float(np.max(np.abs(dt_errs))),
                   note="系统性偏低, 超 declared ±100 ms @ ≥+30 mV"),
    author_GV_EQ_10s=dict(Vh=Vh_auth, k=k_auth),
    verdict="τ 数字化保真(≤2%, +180mV 点 10% 登记); Δt 系统性偏低 0.06-0.19 s 登记并建议换真值表; "
            "G-V 池 25.05 落 psQQ*+E1(23.5) 与 EQ*QQ*Q*(32.8) 之间, 与作者 wt EQ(4-12 mV) 的跨批偏移维持原登记")
json.dump(verdict, open(RES / "判词组件_Kα4_Zenodo10421153_2026-09-26.json", "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
print("tau err %:", np.round(tau_errs, 2))
print("dt err ms:", np.round(dt_errs, 0))
print("saved")
