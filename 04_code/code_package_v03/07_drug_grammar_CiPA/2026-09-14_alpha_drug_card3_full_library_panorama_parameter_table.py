# 2026-09-14_α模型_药物卡3_全药库全景参数表.py
# ============================================================================
# 药物卡3：全药库全景参数表（普查提取卡；判线见
#   预注册_α模型_药物卡3_全药库全景参数表_2026-09-14.md）
#
# 数据：本地 a6k5t\hERG* 全部数据集 subtracted\ted\ted.xlsx（不碰大 csv）
# 单元 = (细胞 × 浓度段)；k_obs = 卡2 锚定非线性估计器（单指数+漂移）
# 门：可判门槛 / L1 Hill 锚比∈[0.7,1.4] / 端点 Welch t / E-4031 R_E 参考
# 合成自检：单浓度型+累积型各一，IC50 找回比∈[0.7,1.4] 且端点 t 可证
#
# 运行：Spyder %runfile '...py' --wdir 。SMOKE=1 冒烟 3 数据集。
# ============================================================================
import os, json
import numpy as np
import openpyxl
from scipy.optimize import curve_fit
from scipy import stats as sst

HERE = os.path.dirname(os.path.abspath(__file__))
PACK = ("D:/Kimi_Agent_细胞仿真工具包扩展以及具身智能20260911/04_细胞线4/结果/"
        "hERG本地运行包_历史版本/hERG本地运行包_v1.1/本地运行包/a6k5t")
SMOKE = os.environ.get("SMOKE", "0") == "1"
SMOKE_SETS = [("a6k5t-osfstorage-hERG-archivelab3", "dofetilide"),
              ("a6k5t-osfstorage-hERG-archivelab4", "pimozide"),
              ("a6k5t-osfstorage-hERG_Phase_II-archivelab1", "moxifloxacin")]

# ---- 判线（钉死）----
MIN_UNITS, MIN_CONCS, MIN_CELLS = 6, 3, 4
MIN_ACTRL_PA, MIN_BSS, MIN_BSS_E, MIN_SWEEPS = 50.0, 0.05, 0.5, 10
RATIO_LINE = (0.7, 1.4)
B_P = 0.01
rng = np.random.default_rng(20260914)

KGRID = np.logspace(-4.5, 0, 26)


def null_fit(t, b):
    """卡2 估计器：b_ss 锚定=末5扫中位；k 网格+d 线性 LS + curve_fit 精修。"""
    bss = float(np.median(b[-5:]))
    tt2 = float(np.sum(t * t))
    E = 1.0 - np.exp(-np.outer(KGRID, t))
    U = bss * E
    d_hat = (float(np.sum(b * t)) - U @ t) / tt2
    R = b - U - np.outer(d_hat, t)
    j = int(np.argmin(np.sum(R * R, axis=1)))
    k0, d0 = float(KGRID[j]), float(d_hat[j])
    try:
        def f(tt, k, d):
            return bss * (1.0 - np.exp(-k * tt)) + d * tt
        p, _ = curve_fit(f, t, b, p0=[k0, d0],
                         bounds=([1e-5, -0.1], [2.0, 0.1]), maxfev=10000)
        return float(p[0]), bss
    except Exception:
        return k0, bss


def conc_nM(conc, unit):
    u = (unit or "").strip().lower()
    return conc * 1000.0 if (u.startswith("u") or u.startswith("µ")) else conc


def load_dataset(tedx):
    wb = openpyxl.load_workbook(tedx, read_only=True)
    rw = [list(r) for r in wb["ResultsWide"].iter_rows(values_only=True)]
    wb.close()
    h = rw[0]
    ix = {k: h.index(k) for k in ("CELLID", "TRACENUM", "ELTIME", "LIQUID", "CONC", "CONCU", "Ramp", "CTLFL")}
    cells = {}
    for r in rw[1:]:
        if r[ix["CELLID"]] is None or r[ix["Ramp"]] is None:
            continue
        try:
            et = float(r[ix["ELTIME"]]) / 1000.0
            rv = float(r[ix["Ramp"]])
            liq = str(r[ix["LIQUID"]]).lower()
            cc = float(r[ix["CONC"]]) if r[ix["CONC"]] is not None else 0.0
            cu = str(r[ix["CONCU"]] or "")
            ctl = str(r[ix["CTLFL"]] or "")
        except (TypeError, ValueError):
            continue
        cells.setdefault(str(r[ix["CELLID"]]), []).append((et, liq, cc, cu, rv, ctl))
    return cells


def extract_units(rows):
    """一个细胞 → (对照A, [(conc_nM, t, b)], e4031)。"""
    tn0 = sorted((r for r in rows
                  if (r[5] == "Y" and (r[2] == 0.0 or "control" in r[1] or "vehicle" in r[1]))
                  or "control" in r[1] or "vehicle" in r[1]), key=lambda r: r[0])
    e40 = sorted((r for r in rows if r[1].startswith("e-4031")), key=lambda r: r[0])
    drg = sorted((r for r in rows if not r[1].startswith("e-4031")
                  and "control" not in r[1] and "vehicle" not in r[1]), key=lambda r: r[0])
    if not tn0 or not drg:
        return None
    A_ctrl = float(np.median([r[4] for r in tn0[-5:]]))
    # 单位归一（v2 钉死）：不同实验室 Ramp 游标单位混杂（pA / nA / A 实证皆存在）
    # |A|<1e-6 → 安培 ×1e12；[1e-6,0.02) → fA 级不可能，排除；[0.02,50) → nA ×1e3；≥50 → pA
    a0 = abs(A_ctrl)
    if a0 < 1e-6:
        scale = 1e12
    elif a0 < 0.02:
        return None
    elif a0 < 50.0:
        scale = 1e3
    else:
        scale = 1.0
    if scale != 1.0:
        rows = [(r[0], r[1], r[2], r[3], r[4] * scale, r[5]) for r in rows]
        tn0 = sorted((r for r in rows
                      if (r[5] == "Y" and (r[2] == 0.0 or "control" in r[1] or "vehicle" in r[1]))
                      or "control" in r[1] or "vehicle" in r[1]), key=lambda r: r[0])
        e40 = sorted((r for r in rows if r[1].startswith("e-4031")), key=lambda r: r[0])
        drg = sorted((r for r in rows if not r[1].startswith("e-4031")
                      and "control" not in r[1] and "vehicle" not in r[1]), key=lambda r: r[0])
        A_ctrl = float(np.median([r[4] for r in tn0[-5:]]))
    if abs(A_ctrl) < MIN_ACTRL_PA:
        return None
    segs = []
    by_c = {}
    for r in drg:
        by_c.setdefault(conc_nM(r[2], r[3]), []).append(r)
    for C, rs in sorted(by_c.items()):
        t = np.array([x[0] - rs[0][0] for x in rs])
        b = 1.0 - np.array([x[4] for x in rs]) / A_ctrl
        segs.append((C, t, b))
    out = {"A_ctrl": A_ctrl, "segs": segs}
    if len(e40) >= MIN_SWEEPS:
        A_ss = float(np.median([r[4] for r in drg[-5:]]))
        if abs(A_ss) > 20:
            te = np.array([x[0] - e40[0][0] for x in e40])
            be = 1.0 - np.array([x[4] for x in e40]) / A_ss
            out["e4031"] = (te, be)
    return out


def hill_fit(Cs, bs):
    def hill(C, ic50, n):
        return C ** n / (ic50 ** n + C ** n)
    C = np.array(Cs, float); b = np.array(bs, float)
    p0 = [float(np.exp(np.mean(np.log(np.clip(C, 1e-9, None))))), 1.0]
    p, _ = curve_fit(hill, C, b, p0=p0, bounds=([1e-3, 0.2], [1e8, 4.0]), maxfev=20000)
    return float(p[0]), float(p[1])


def lab_anchor(ds_dir):
    fx = os.path.join(ds_dir, "subtracted", "tables", "concentration-inhibition.xlsx")
    if not os.path.exists(fx):
        return None
    wb = openpyxl.load_workbook(fx, read_only=True)
    if "ic50nh" not in wb.sheetnames:
        wb.close(); return None
    rows = [list(r) for r in wb["ic50nh"].iter_rows(values_only=True)]
    wb.close()
    hdr = [str(x) for x in rows[0]]
    r = rows[1]
    ic50 = float(r[hdr.index("IC50")])
    nh = float(r[hdr.index("nh")])
    unit = str(r[hdr.index("CONCU")])
    return dict(ic50_nM=conc_nM(ic50, unit), nh=nh)


def analyze_dataset(arch, drug, ds_dir):
    cells = load_dataset(os.path.join(ds_dir, "subtracted", "ted", "ted.xlsx"))
    units = []
    kE = []
    for cid, rows in sorted(cells.items()):
        ex = extract_units(rows)
        if not ex:
            continue
        for C, t, b in ex["segs"]:
            if len(t) < MIN_SWEEPS:
                continue
            bss = float(np.median(b[-5:]))
            k_obs = None
            if bss >= MIN_BSS:
                k_obs, _ = null_fit(t, b)
            units.append(dict(cell=cid, conc_nM=C, b_ss=bss, k_obs=k_obs, n_sweeps=len(t)))
        if "e4031" in ex:
            te, be = ex["e4031"]
            bss_e = float(np.median(be[-5:]))
            if bss_e >= MIN_BSS_E:
                k_e, _ = null_fit(te, be)
                kE.append(k_e)
    ncells = len(set(u["cell"] for u in units))
    concs = sorted(set(round(u["conc_nM"], 6) for u in units))
    res = dict(n_units=len(units), n_cells=ncells, n_concs=len(concs),
               concs_nM=concs, units=units, kE4031_med=float(np.median(kE)) if kE else None)
    # 可判门槛
    if not (len(units) >= MIN_UNITS and len(concs) >= MIN_CONCS and ncells >= MIN_CELLS):
        res["判"] = "样本不足"
        return res
    # L1
    ic50, nh = hill_fit([u["conc_nM"] for u in units], [u["b_ss"] for u in units])
    anch = lab_anchor(ds_dir)
    L1 = dict(ic50_self=ic50, nh_self=nh)
    if anch:
        L1.update(ic50_lab=anch["ic50_nM"], nh_lab=anch["nh"],
                  ratio_ic50=ic50 / anch["ic50_nM"], ratio_nh=nh / anch["nh"])
        L1["过"] = bool(RATIO_LINE[0] <= L1["ratio_ic50"] <= RATIO_LINE[1]
                        and RATIO_LINE[0] <= L1["ratio_nh"] <= RATIO_LINE[1])
    else:
        L1["过"] = None
        L1["注"] = "实验室锚缺失"
    res["L1"] = L1
    # 浓度依赖（端点 t）
    kk = [(u["conc_nM"], u["k_obs"]) for u in units if u["k_obs"]]
    gB = dict(判="组不足")
    if len(set(round(c, 6) for c, _ in kk)) >= 3:
        C = np.array([c for c, _ in kk]); k = np.array([v for _, v in kk])
        grids = sorted(set(np.round(C, 6)))
        lo = C <= grids[0] + 1e-12
        hi = C >= grids[-2] - 1e-12
        if lo.sum() >= 2 and hi.sum() >= 2:
            tstat, p2 = sst.ttest_ind(k[hi], k[lo], equal_var=False)
            p1 = float(p2 / 2) if tstat > 0 else 1.0
            gB = dict(判=("浓度依赖可证" if (p1 < B_P and k[hi].mean() > k[lo].mean())
                          else "不可证(饱和方向)"),
                      p_one_side=p1, k_lo=float(k[lo].mean()), k_hi=float(k[hi].mean()),
                      n_lo=int(lo.sum()), n_hi=int(hi.sum()))
    res["B"] = gB
    # E-4031 参照
    if kE and gB.get("k_hi"):
        res["R_E"] = float(np.median(kE) / gB["k_hi"])
    res["判"] = "可判"
    return res


# ------------------------------------------------------------ 合成自检
def synth_check():
    def synth_units(kind):
        units = []
        concs = [0.3, 1.0, 3.0, 10.0]
        if kind == "single":
            for ci, C in enumerate(concs):
                for j in range(4):
                    units.append((f"cell_{ci}_{j}", C))
        else:
            for j in range(5):
                for C in concs:
                    units.append((f"cell_{j}", C))
        return units
    ok_all = True
    for kind in ("single", "累积"):
        units = []
        true_ic, true_nh, kon, koff = 2.0, 1.1, 2e-3, 2e-3
        for cid, C in synth_units(kind):
            bss = C ** true_nh / (true_ic ** true_nh + C ** true_nh)
            k = kon * C + koff
            units.append(dict(cell=cid, conc_nM=C, b_ss=bss, k_obs=k + rng.normal(0, 2e-4)))
        ic50, nh = hill_fit([u["conc_nM"] for u in units], [u["b_ss"] for u in units])
        ratio = ic50 / true_ic
        C = np.array([u["conc_nM"] for u in units]); k = np.array([u["k_obs"] for u in units])
        grids = sorted(set(np.round(C, 6)))
        lo = C <= grids[0]; hi = C >= grids[-2]
        tstat, p2 = sst.ttest_ind(k[hi], k[lo], equal_var=False)
        p1 = float(p2 / 2) if tstat > 0 else 1.0
        ok = RATIO_LINE[0] <= ratio <= RATIO_LINE[1] and p1 < B_P
        ok_all &= ok
        print(f"  [合成 {kind}] IC50 找回比 {ratio:.3f} 端点p {p1:.5f} -> {'过' if ok else '不过'}", flush=True)
    return bool(ok_all)


# ------------------------------------------------------------ 主跑
print("=" * 76, flush=True)
print(" 药物卡3 · 全药库全景参数表（普查提取卡）", flush=True)
print("=" * 76, flush=True)

print("\n[合成自检]", flush=True)
synth_ok = synth_check()
print(f"  合成自检 {'过' if synth_ok else '不过——全卡降级登记'}", flush=True)

# 数据集清单
datasets = []
for arch in sorted(os.listdir(PACK)):
    d0 = os.path.join(PACK, arch)
    if not (os.path.isdir(d0) and arch.startswith("a6k5t-osfstorage-hERG")):
        continue
    for drug in sorted(os.listdir(d0)):
        ds = os.path.join(d0, drug)
        if os.path.isdir(ds) and os.path.exists(os.path.join(ds, "subtracted", "ted", "ted.xlsx")):
            datasets.append((arch, drug, ds))
if SMOKE:
    datasets = [d for d in datasets if (d[0], d[1]) in SMOKE_SETS]
print(f"\n数据集: {len(datasets)}（{'冒烟' if SMOKE else '全量'}）", flush=True)

result = {"预注册": "预注册_α模型_药物卡3_全药库全景参数表_2026-09-14.md",
          "合成自检": synth_ok, "datasets": {}, "汇总": {}}
for i, (arch, drug, ds) in enumerate(datasets):
    key = f"{arch.replace('a6k5t-osfstorage-', '')}|{drug}"
    try:
        r = analyze_dataset(arch, drug, ds)
    except Exception as e:
        r = dict(判="提取异常", err=str(e)[:200])
    result["datasets"][key] = r
    tag = r.get("判", "?")
    l1 = r.get("L1") or {}
    line = (f"  [{i + 1}/{len(datasets)}] {key}: {tag} 单元{r.get('n_units', 0)}"
            f" IC50比 {l1.get('ratio_ic50', float('nan')):.2f}" if l1 else
            f"  [{i + 1}/{len(datasets)}] {key}: {tag} 单元{r.get('n_units', 0)}")
    print(line, flush=True)

# 汇总（只描述不判决）
rows = []
for key, r in result["datasets"].items():
    if r.get("判") != "可判":
        continue
    drug = key.split("|")[1]
    l1 = r.get("L1") or {}
    rows.append(dict(ds=key, drug=drug, n_units=r["n_units"],
                     ic50_ratio=l1.get("ratio_ic50"), nh_ratio=l1.get("ratio_nh"),
                     L1过=l1.get("过"), B判=(r.get("B") or {}).get("判"), R_E=r.get("R_E")))
by_drug = {}
for r in rows:
    if r["ic50_ratio"]:
        by_drug.setdefault(r["drug"], []).append(r["ic50_ratio"])
summ = {d: dict(n=len(v), ratio_med=float(np.median(v)),
                ratio_span=float(max(v) / min(v))) for d, v in by_drug.items()}
result["汇总"] = dict(可判数据集=len(rows), L1过率=float(np.mean([bool(r["L1过"]) for r in rows])) if rows else None,
                    浓度依赖可证率=float(np.mean([r["B判"] == "浓度依赖可证" for r in rows])) if rows else None,
                    分药IC50比=summ)

print("\n" + "=" * 76, flush=True)
print(f" 可判 {len(rows)}/{len(datasets)}；L1过率 {result['汇总']['L1过率']}; "
      f"浓度依赖可证率 {result['汇总']['浓度依赖可证率']}", flush=True)

sfx = "_冒烟" if SMOKE else ""
fj = os.path.join(HERE, f"2026-09-14_α模型_药物卡3_全药库全景参数表{sfx}_结果.json")


def _clean(o):
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, (np.bool_,)):
        return bool(o)
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, np.floating):
        return float(o) if np.isfinite(o) else None
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, float) and not np.isfinite(o):
        return None
    return o


json.dump(_clean(result), open(fj, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("  结果落盘:", fj, flush=True)

# CSV 全景表
import csv as _csv
fc = os.path.join(HERE, f"2026-09-14_α模型_药物卡3_全景表{sfx}.csv")
with open(fc, "w", newline="", encoding="utf-8-sig") as w:
    wr = _csv.DictWriter(w, fieldnames=["ds", "drug", "n_units", "ic50_ratio", "nh_ratio", "L1过", "B判", "R_E"])
    wr.writeheader()
    for r in rows:
        wr.writerow(r)
print("  全景表落盘:", fc, flush=True)

# ------------------------------------------------------------ 图
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

fig, ax = plt.subplots(1, 3, figsize=(15, 5))
xs = [r["ic50_ratio"] for r in rows if r["ic50_ratio"]]
ax[0].hist(xs, bins=20, color="#4472C4")
ax[0].axvline(0.7, ls="--", c="r"); ax[0].axvline(1.4, ls="--", c="r")
ax[0].set_title(f"L1 IC50 自提/锚 比分布（n={len(xs)}）"); ax[0].grid(alpha=0.3)
ys = [r["R_E"] for r in rows if r["R_E"]]
ax[1].hist(np.log10(np.clip(ys, 0.05, 1e4)), bins=20, color="#70AD47")
ax[1].axvline(np.log10(3), ls="--", c="r"); ax[1].axvline(np.log10(2), ls="--", c="orange")
ax[1].set_title(f"log10 R_E 分布（n={len(ys)}，红=灌注否线3）"); ax[1].grid(alpha=0.3)
bd = {}
for r in rows:
    bd.setdefault(r["drug"], []).append(r["ic50_ratio"])
names = sorted(bd, key=lambda d: np.median([x for x in bd[d] if x]))
meds = [np.median([x for x in bd[d] if x]) for d in names]
spans = [(min([x for x in bd[d] if x]), max([x for x in bd[d] if x])) for d in names]
yv = np.arange(len(names))
ax[2].hlines(yv, [s[0] for s in spans], [s[1] for s in spans], color="#999")
ax[2].plot(meds, yv, "o", color="#C00000")
ax[2].axvline(0.7, ls="--", c="r"); ax[2].axvline(1.4, ls="--", c="r")
ax[2].set_yticks(yv); ax[2].set_yticklabels(names, fontsize=6)
ax[2].set_title("分药 IC50 比（点=中位，线=跨实验室极差）"); ax[2].grid(alpha=0.3)
fig.tight_layout()
fp = os.path.join(HERE, f"2026-09-14_α模型_药物卡3_全景参数表{sfx}.png")
fig.savefig(fp, dpi=130, bbox_inches="tight")
print("  图落盘:", fp, flush=True)
